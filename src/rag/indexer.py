# src/rag/indexer.py
from pathlib import Path
from typing import List
import ollama
import chromadb
from chromadb.config import Settings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader


DOCS_DIR = Path("rag_docs")
DOCS_DIR.mkdir(exist_ok=True)

CHROMA_DIR = Path("rag_chroma")
CHROMA_DIR.mkdir(exist_ok=True)

EMBED_MODEL = "nomic-embed-text"

_client = chromadb.PersistentClient(
    path=str(CHROMA_DIR),
    settings=Settings(anonymized_telemetry=False),
)


# ============================================================
# Lectura de diferentes tipos de archivo
# ============================================================

def _leer_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    texto = []
    for page in reader.pages:
        texto.append(page.extract_text() or "")
    return "\n".join(texto)


def _leer_texto(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _leer_csv(path: Path) -> str:
    """Lee un CSV y devuelve como tabla markdown simple."""
    import csv
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            lector = csv.reader(f)
            lineas = list(lector)
    except Exception:
        return _leer_texto(path)

    if not lineas:
        return ""

    resultado = []
    resultado.append(" | ".join(lineas[0]))
    resultado.append(" | ".join(["---"] * len(lineas[0])))
    for fila in lineas[1:]:
        resultado.append(" | ".join(fila))
    return "\n".join(resultado)


def _leer_json(path: Path) -> str:
    """Lee un JSON y lo formatea bonito."""
    import json
    try:
        data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
        return json.dumps(data, indent=2, ensure_ascii=False)
    except Exception:
        return _leer_texto(path)


def _leer_excel(path: Path) -> str:
    """Lee un Excel y devuelve el contenido de todas las hojas."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "(openpyxl no instalado, no se puede leer el Excel)"

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        resultado = []
        for hoja in wb.sheetnames:
            ws = wb[hoja]
            resultado.append(f"\n=== HOJA: {hoja} ===\n")
            for fila in ws.iter_rows(values_only=True):
                resultado.append(
                    " | ".join(str(c) if c is not None else "" for c in fila)
                )
        wb.close()
        return "\n".join(resultado)
    except Exception as e:
        return f"(Error leyendo Excel: {e})"


def _leer_docx(path: Path) -> str:
    """Lee un Word."""
    try:
        from docx import Document
    except ImportError:
        return "(python-docx no instalado, no se puede leer el .docx)"

    try:
        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        return f"(Error leyendo docx: {e})"


def leer_documento(path: Path) -> str:
    """Lee cualquier documento soportado."""
    ext = path.suffix.lower()

    if ext == ".pdf":
        return _leer_pdf(path)
    if ext == ".csv":
        return _leer_csv(path)
    if ext == ".json":
        return _leer_json(path)
    if ext in (".xlsx", ".xls"):
        return _leer_excel(path)
    if ext == ".docx":
        return _leer_docx(path)

    # Código, texto, configs, etc → leer como texto
    return _leer_texto(path)


# ============================================================
# Chunking y embeddings
# ============================================================

def _chunkear(texto: str) -> List[str]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=150,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    return splitter.split_text(texto)


def _embed(textos: List[str]) -> List[List[float]]:
    vectores = []
    for t in textos:
        resp = ollama.embeddings(model=EMBED_MODEL, prompt=t)
        vectores.append(resp["embedding"])
    return vectores


# ============================================================
# Indexar / listar / eliminar
# ============================================================

def indexar_documento(project_id: str, nombre_archivo: str, contenido: bytes) -> dict:
    proyecto_dir = DOCS_DIR / project_id
    proyecto_dir.mkdir(exist_ok=True)
    path = proyecto_dir / nombre_archivo
    path.write_bytes(contenido)

    texto = leer_documento(path)
    if not texto.strip():
        return {"error": "El documento no contiene texto extraíble"}

    chunks = _chunkear(texto)
    vectores = _embed(chunks)

    coleccion_nombre = f"proj_{project_id}"
    coleccion = _client.get_or_create_collection(name=coleccion_nombre)

    ids = [f"{nombre_archivo}_{i}" for i in range(len(chunks))]
    metadatas = [{"archivo": nombre_archivo, "chunk": i} for i in range(len(chunks))]

    coleccion.add(
        ids=ids,
        embeddings=vectores,
        documents=chunks,
        metadatas=metadatas,
    )

    return {
        "archivo": nombre_archivo,
        "chunks": len(chunks),
        "caracteres": len(texto),
    }


def listar_documentos(project_id: str) -> List[dict]:
    proyecto_dir = DOCS_DIR / project_id
    if not proyecto_dir.exists():
        return []
    return [
        {"nombre": f.name, "size": f.stat().st_size}
        for f in sorted(proyecto_dir.iterdir())
        if f.is_file()
    ]


def eliminar_documento(project_id: str, nombre_archivo: str) -> bool:
    path = DOCS_DIR / project_id / nombre_archivo
    if path.exists():
        path.unlink()

    coleccion_nombre = f"proj_{project_id}"
    try:
        coleccion = _client.get_collection(name=coleccion_nombre)
        resultados = coleccion.get(where={"archivo": nombre_archivo})
        if resultados["ids"]:
            coleccion.delete(ids=resultados["ids"])
    except Exception:
        pass

    return True
