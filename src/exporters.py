# src/exporters.py
"""Exportadores multi-formato para respuestas de NIAH."""
import io
import re
import json as _json
from datetime import datetime


def _parse_bloques(md: str) -> list[dict]:
    bloques, lineas, i = [], md.split("\n"), 0
    while i < len(lineas):
        l = lineas[i]
        if l.startswith("### "):
            bloques.append({"t": "h3", "v": l[4:].strip()})
        elif l.startswith("## "):
            bloques.append({"t": "h2", "v": l[3:].strip()})
        elif l.startswith("# "):
            bloques.append({"t": "h1", "v": l[2:].strip()})
        elif re.match(r"^\s*[\-\*]\s+", l):
            bloques.append({"t": "li", "v": re.sub(r"^\s*[\-\*]\s+", "", l)})
        elif l.strip().startswith("|") and l.strip().endswith("|"):
            filas = []
            while i < len(lineas) and lineas[i].strip().startswith("|"):
                f = lineas[i].strip()
                if not re.match(r"^\|[\s\-:|]+\|$", f):
                    filas.append([c.strip() for c in f.strip("|").split("|")])
                i += 1
            bloques.append({"t": "tabla", "v": filas})
            continue
        elif l.strip():
            bloques.append({"t": "p", "v": l})
        i += 1
    return bloques


def _strip_inline(s: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"\*(.+?)\*", r"\1", s)
    s = re.sub(r"`(.+?)`", r"\1", s)
    s = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1 (\2)", s)
    return s


def _plain(md: str) -> str:
    out = []
    for b in _parse_bloques(md):
        t, v = b["t"], b["v"]
        if t == "tabla":
            for fila in v:
                out.append(" | ".join(fila))
        elif t in ("h1", "h2", "h3"):
            out.append(_strip_inline(v))
            out.append("")
        elif t == "li":
            out.append("• " + _strip_inline(v))
        else:
            out.append(_strip_inline(v))
    return "\n".join(out)
def to_txt(titulo: str, contenido: str):
    txt = f"{titulo}\n{'='*len(titulo)}\n\n{_plain(contenido)}\n"
    return txt.encode("utf-8"), "text/plain; charset=utf-8", "txt"


def to_md(titulo: str, contenido: str):
    md = f"# {titulo}\n\n{contenido}\n"
    return md.encode("utf-8"), "text/markdown; charset=utf-8", "md"


def to_json(titulo: str, contenido: str):
    data = {
        "titulo": titulo,
        "contenido": contenido,
        "contenido_plano": _plain(contenido),
        "generado": datetime.now().isoformat(),
    }
    return _json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8"), "application/json", "json"


def to_html(titulo: str, contenido: str):
    partes = []
    for b in _parse_bloques(contenido):
        t, v = b["t"], b["v"]
        if t == "h1":
            partes.append(f"<h1>{_strip_inline(v)}</h1>")
        elif t == "h2":
            partes.append(f"<h2>{_strip_inline(v)}</h2>")
        elif t == "h3":
            partes.append(f"<h3>{_strip_inline(v)}</h3>")
        elif t == "li":
            partes.append(f"<li>{_strip_inline(v)}</li>")
        elif t == "tabla":
            rows = "".join("<tr>" + "".join(f"<td>{c}</td>" for c in f) + "</tr>" for f in v)
            partes.append(f"<table>{rows}</table>")
        else:
            partes.append(f"<p>{_strip_inline(v)}</p>")
    body = "\n".join(partes)
    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8"><title>{titulo}</title>
<style>body{{font-family:system-ui,sans-serif;max-width:900px;margin:40px auto;padding:20px;color:#1d1d1f}}
h1{{color:#0a84ff}}h2{{color:#0a84ff;margin-top:32px}}h3{{color:#334155}}
table{{border-collapse:collapse;width:100%;margin:12px 0}}
td{{border:1px solid #ddd;padding:6px 10px}}li{{margin:4px 0}}</style>
</head><body><h1>{titulo}</h1>{body}</body></html>"""
    return html.encode("utf-8"), "text/html; charset=utf-8", "html"


def to_pdf(titulo: str, contenido: str):
    from weasyprint import HTML
    html_bytes, _, _ = to_html(titulo, contenido)
    pdf = HTML(string=html_bytes.decode("utf-8")).write_pdf()
    return pdf, "application/pdf", "pdf"
def to_docx(titulo: str, contenido: str):
    from docx import Document
    from docx.shared import RGBColor
    doc = Document()
    h = doc.add_heading(titulo, level=0)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x0A, 0x84, 0xFF)
    for b in _parse_bloques(contenido):
        t, v = b["t"], b["v"]
        if t == "h1":
            doc.add_heading(_strip_inline(v), level=1)
        elif t == "h2":
            doc.add_heading(_strip_inline(v), level=2)
        elif t == "h3":
            doc.add_heading(_strip_inline(v), level=3)
        elif t == "li":
            doc.add_paragraph(_strip_inline(v), style="List Bullet")
        elif t == "tabla" and v:
            ncols = max(len(f) for f in v)
            tbl = doc.add_table(rows=len(v), cols=ncols)
            tbl.style = "Light Grid Accent 1"
            for ri, fila in enumerate(v):
                for ci, cell in enumerate(fila):
                    if ci < ncols:
                        tbl.rows[ri].cells[ci].text = _strip_inline(cell)
        elif v.strip():
            doc.add_paragraph(_strip_inline(v))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"


def to_xlsx(titulo: str, contenido: str):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Respuesta"
    ws["A1"] = titulo
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="0A84FF")
    ws.merge_cells("A1:C1")
    fila = 3
    for b in _parse_bloques(contenido):
        t, v = b["t"], b["v"]
        if t == "tabla" and v:
            for fila_data in v:
                for ci, cell in enumerate(fila_data, start=1):
                    ws.cell(row=fila, column=ci, value=_strip_inline(cell))
                fila += 1
            fila += 1
        elif t in ("h1", "h2", "h3"):
            c = ws.cell(row=fila, column=1, value=_strip_inline(v))
            c.font = Font(bold=True, size=12 if t == "h3" else 14)
            fila += 1
        elif t == "li":
            ws.cell(row=fila, column=1, value="• " + _strip_inline(v))
            fila += 1
        elif v.strip():
            ws.cell(row=fila, column=1, value=_strip_inline(v))
            fila += 1
    ws.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
def to_pptx(titulo: str, contenido: str):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor as PPColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # Portada
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fondo = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    fondo.fill.solid()
    fondo.fill.fore_color.rgb = PPColor(0x0A, 0x84, 0xFF)
    fondo.line.fill.background()
    tb = s.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(2))
    tf = tb.text_frame
    tf.text = titulo
    p = tf.paragraphs[0]
    p.font.size = Pt(44); p.font.bold = True; p.font.color.rgb = PPColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER
    tb2 = s.shapes.add_textbox(Inches(1), Inches(4.7), Inches(11.33), Inches(1))
    tf2 = tb2.text_frame
    tf2.text = f"Generado el {datetime.now().strftime('%d/%m/%Y a las %H:%M')}"
    p2 = tf2.paragraphs[0]
    p2.font.size = Pt(18); p2.font.color.rgb = PPColor(0xE6, 0xF1, 0xFF)
    p2.alignment = PP_ALIGN.CENTER

    def nuevo_slide(header=None):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        barra = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.3), prs.slide_height)
        barra.fill.solid(); barra.fill.fore_color.rgb = PPColor(0x34, 0xC7, 0x59); barra.line.fill.background()
        y_top = Inches(1.6) if header else Inches(0.4)
        box = sl.shapes.add_textbox(Inches(0.7), y_top, Inches(12), Inches(5.5))
        tfc = box.text_frame
        tfc.word_wrap = True
        if header:
            tb_t = sl.shapes.add_textbox(Inches(0.7), Inches(0.4), Inches(12), Inches(1))
            tf_t = tb_t.text_frame
            tf_t.text = header
            pt = tf_t.paragraphs[0]
            pt.font.size = Pt(30); pt.font.bold = True; pt.font.color.rgb = PPColor(0x0A, 0x84, 0xFF)
        return tfc

    bloques = _parse_bloques(contenido)
    hay_headers = any(b["t"] in ("h1", "h2") for b in bloques)

    if hay_headers:
        actual_tf = None
        primera = True
        for b in bloques:
            t, v = b["t"], b["v"]
            if t in ("h1", "h2"):
                actual_tf = nuevo_slide(header=_strip_inline(v))
                primera = True
            elif actual_tf is not None:
                p = actual_tf.paragraphs[0] if primera else actual_tf.add_paragraph()
                primera = False
                if t == "li":
                    p.text = "• " + _strip_inline(v)
                else:
                    p.text = _strip_inline(v)
                p.font.size = Pt(20 if t == "h3" else 16)
                if t == "h3":
                    p.font.bold = True
    else:
        actual_tf = nuevo_slide(header="Contenido")
        primera = True
        contador = 0
        for b in bloques:
            if b["t"] == "tabla":
                texto = " | ".join(b["v"][0]) if b["v"] else ""
            else:
                texto = _strip_inline(b["v"])
            if not texto.strip():
                continue
            p = actual_tf.paragraphs[0] if primera else actual_tf.add_paragraph()
            primera = False
            p.text = ("• " if b["t"] == "li" else "") + texto
            p.font.size = Pt(16)
            contador += 1
            if contador >= 8:
                actual_tf = nuevo_slide(header="Continuación")
                primera = True
                contador = 0

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx"


EXPORTERS = {
    "txt": to_txt,
    "md": to_md,
    "json": to_json,
    "html": to_html,
    "pdf": to_pdf,
    "docx": to_docx,
    "xlsx": to_xlsx,
    "pptx": to_pptx,
}


def exportar(formato: str, titulo: str, contenido: str):
    fn = EXPORTERS.get(formato.lower())
    if not fn:
        raise ValueError(f"Formato no soportado: {formato}")
    return fn(titulo, contenido)

